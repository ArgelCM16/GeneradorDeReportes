# -*- coding: utf-8 -*-
"""
Revisa los .docx que genera la prueba "word" (casos/word.html).

Siempre revisa con la biblioteca estándar: que el ZIP esté bien, que cada XML
sea válido y que estén todas las partes. Si está instalado python-docx
(pip install python-docx), además abre los documentos y revisa estilos,
textos, citas, listas, imágenes, tabla, tamaño de hoja y márgenes.
"""
import base64
import io
import os
import re
import zipfile
import xml.dom.minidom

PARTS = ['[Content_Types].xml', '_rels/.rels', 'word/document.xml', 'word/styles.xml', 'word/numbering.xml',
         'word/settings.xml', 'word/footer1.xml', 'word/_rels/document.xml.rels', 'docProps/core.xml']


def check_docx_files(files, work_dir):
    results = []

    def ok(name, cond):
        results.append(('PASS ' if cond else 'FAIL ') + name)

    if not files:
        ok('word: la prueba no entregó archivos', False)
    try:
        import docx  # noqa: F401
        have_docx = True
    except ImportError:
        have_docx = False
        results.append('PASS word: python-docx no está instalado; solo se revisa el ZIP y el XML')

    for name, b64 in files.items():
        data = base64.b64decode(b64)
        z = zipfile.ZipFile(io.BytesIO(data))
        ok(f'{name}: ZIP íntegro', z.testzip() is None)
        if name == 'zip.zip':
            ok('zip: contenido', z.read('a.txt') == b'hola' and z.read('ñ/b.txt') == bytes([1, 2, 3]))
            continue
        bad = []
        for item in z.namelist():
            if item.endswith('.xml') or item.endswith('.rels'):
                try:
                    xml.dom.minidom.parseString(z.read(item))
                except Exception as e:
                    bad.append(f'{item}: {e}')
        ok(f'{name}: XML válido {bad}', not bad)
        ok(f'{name}: partes completas', all(p in z.namelist() for p in PARTS))
        docxml = z.read('word/document.xml').decode('utf-8')
        styles = z.read('word/styles.xml')
        ok(f'{name}: índice como campo de Word', 'TOC \\o "1-2"' in docxml and docxml.count('fldCharType="begin"') == 1)
        ok(f'{name}: código con colores', '<w:color w:val="CF222E"/>' in docxml)

        if name == 'normal.docx':
            ok('normal: footer con número de página', b'PAGE' in z.read('word/footer1.xml') and 'footerReference' in docxml)
            ok('normal: Georgia', b'w:ascii="Georgia"' in styles)
            ok('normal: justificado', 'w:jc w:val="both"' in docxml)
        elif name == 'cover.docx':
            ok('portada: sin números de página', 'footerReference' not in docxml)
            ok('portada: primera hoja distinta', '<w:titlePg/>' in docxml)
            ok('portada: Times y doble espacio', b'Times New Roman' in styles and b'w:line="480"' in styles)

        if not have_docx:
            continue
        import docx
        from docx.shared import Emu
        path = os.path.join(work_dir, 'revision_' + name)
        with open(path, 'wb') as f:
            f.write(data)
        d = docx.Document(path)
        texts = [p.text for p in d.paragraphs]
        alltext = '\n'.join(texts)
        names = [p.style.name for p in d.paragraphs]
        ok(f'{name}: título como Heading 1', any(t == 'Práctica de Ohm & <otros>' and s == 'Heading 1' for t, s in zip(texts, names)))
        ok(f'{name}: subtítulo como Heading 2', any(t == 'Introducción' and s == 'Heading 2' for t, s in zip(texts, names)))
        para = next((p for p in d.paragraphs if p.text.startswith('Texto con')), None)
        runs = {r.text: (r.bold, r.italic, r.underline) for r in para.runs} if para else {}
        ok(f'{name}: negritas, cursivas y subrayado', runs.get('negritas', (0,))[0] is True and runs.get('cursivas', (0, 0))[1] is True and runs.get('subrayado', (0, 0, 0))[2] is True)
        cite = '[1]' if name == 'normal.docx' else '(Pérez, 2020)'
        ok(f'{name}: cita {cite}', para is not None and cite in para.text)
        numbered = [p for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
        ok(f'{name}: elementos de lista ({len(numbered)})', len(numbered) == 5)
        ok(f'{name}: imágenes ({len(d.inline_shapes)})', len(d.inline_shapes) == 3)
        sec = d.sections[0]
        width = sec.page_width - sec.left_margin - sec.right_margin
        ok(f'{name}: imágenes caben a lo ancho', all(s.width <= width + 10 for s in d.inline_shapes))
        ok(f'{name}: descripciones', 'Figura 1: Diagrama PNG' in alltext and 'Tabla 1: Datos medidos' in alltext)
        ok(f'{name}: tabla', len(d.tables) == 1 and [c.text for c in d.tables[0].rows[1].cells] == ['5', '10', '0.5'])
        if name == 'normal.docx':
            ok('normal: hoja carta', sec.page_width == Emu(12240 * 635) and sec.page_height == Emu(15840 * 635))
            ok('normal: márgenes de 2 cm', abs(sec.left_margin.cm - 2) < 0.01)
            ok('normal: datos del encabezado', 'Alumno: Ana López | Matrícula: 2109045 | Grupo: IDY-7A' in alltext)
        elif name == 'cover.docx':
            ok('portada: hoja oficio', sec.page_height == Emu(20160 * 635))
            ok('portada: nombre de la tarea e integrantes', 'Ley de Ohm' in texts and 'Ana López (1)' in texts)

    passed = sum(1 for r in results if r.startswith('PASS'))
    failed = [r for r in results if r.startswith('FAIL')]
    return passed, failed
